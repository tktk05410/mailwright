"""索引持久化与文档读取。"""

import json
from pathlib import Path

import jieba
import numpy as np
from rank_bm25 import BM25Okapi

from core.models import Chunk
from preprocess.chunker import chunk_handbook
from rag.indexer import KnowledgeIndex

DOC_EXTS = (".md", ".txt", ".pdf", ".docx")
# 管线版本，改动后缓存失效
PIPELINE_VERSION = 2


def read_document(path):
    """读取文档为纯文本。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        try:
            return path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            try:
                return path.read_text(encoding="gbk")
            except UnicodeDecodeError as ex:
                raise ValueError(f"无法解码文件（请另存为 UTF-8）：{path.name}") from ex
    if suffix == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            return "\n".join(page.get_text() for page in doc)
    if suffix == ".docx":
        from docx import Document
        d = Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" ｜ ".join(cells))
        return "\n".join(parts)
    raise ValueError(f"暂不支持的格式：{suffix}")


def load_handbook_chunks(folder, exclude=(), errors=None):
    """读取手册目录。"""
    folder = Path(folder)
    chunks = []
    if not folder.exists():
        return chunks
    for f in sorted(folder.iterdir()):
        if not (f.is_file() and f.suffix.lower() in DOC_EXTS and f.name not in exclude):
            continue
        try:
            text = read_document(f)
        except Exception as ex:
            if errors is not None:
                errors.append((f.name, str(ex)))
            continue
        if not text.strip():
            if errors is not None:
                errors.append((f.name, "没有可提取的文字（扫描件？）"))
            continue
        chunks.extend(chunk_handbook(text, f.stem))
    return chunks


def load_style_chunks(path, prefix="st"):
    """读取语气样本。"""
    p = Path(path)
    if not p.exists():
        return []
    try:
        items = json.loads(p.read_text(encoding="utf-8"))
        if not isinstance(items, list):
            raise ValueError("数据不是列表")
    except Exception:
        return []
    return [
        Chunk(
            chunk_id=f"{prefix}_{i}",
            text=it.get("body", ""),
            kind="style",
            metadata={
                "to": it.get("to", ""),
                "subject": it.get("subject", ""),
                "date": it.get("date", ""),  # 用于时间衰减
            },
        )
        for i, it in enumerate(items, 1)
        if isinstance(it, dict) and it.get("body", "")
    ]


def sources_snapshot(folder, extra_files=()):
    """源文件指纹。"""
    snap = {}
    folder = Path(folder)
    if folder.exists():
        for f in sorted(folder.iterdir()):
            if f.is_file():
                snap[f.name] = [round(f.stat().st_mtime, 2), f.stat().st_size]
    for f in extra_files:
        p = Path(f)
        if p.exists():
            snap[p.name] = [round(p.stat().st_mtime, 2), p.stat().st_size]
    return snap


def save_index(index, folder, sources=None):
    """先写向量再写清单。"""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    np.savez(folder / "vectors.npz", **index.vec)
    payload = {
        "pipeline": PIPELINE_VERSION,
        "signature": index.embedder.signature,
        "sources": sources or {},
        "kinds": {kind: [c.__dict__ for c in group] for kind, group in index.order.items()},
    }
    (folder / "chunks.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8"
    )


def load_index(folder, embedder=None, sources=None):
    """从磁盘恢复索引。"""
    folder = Path(folder)
    if not (folder / "chunks.json").exists():
        return None
    try:
        payload = json.loads((folder / "chunks.json").read_text(encoding="utf-8"))
    except Exception:
        return None

    index = KnowledgeIndex(embedder)
    # 三道一致性校验
    if payload.get("pipeline") != PIPELINE_VERSION:
        return None
    if payload.get("signature") != index.embedder.signature:
        return None
    if payload.get("sources") != (sources or {}):
        return None

    try:
        vectors = np.load(folder / "vectors.npz")
        for kind, items in payload["kinds"].items():
            group = [Chunk(**it) for it in items]
            arr = np.asarray(vectors[kind], dtype=np.float32)
            # 行数必须与块数一致
            if arr.ndim != 2 or arr.shape[0] != len(group):
                return None
            index.order[kind] = group
            index.chunks.extend(group)
            index.bm25[kind] = BM25Okapi([list(jieba.cut_for_search(c.text)) for c in group])
            index.vec[kind] = arr
    except Exception:
        return None
    return index
